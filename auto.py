
from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

log = logging.getLogger("auto")

DEFAULT_INPUT = "input.txt"
DEFAULT_OUTPUT = "output.xlsx"
DEFAULT_SHEET = "Data"
INVALID_SHEET_CHARS = set("[]:*?/\\")
MAX_SHEET_NAME_LEN = 31  # Excel limit


def parse_records(lines: Iterable[str]) -> list[dict[str, str]]:
    
    records: list[dict[str, str]] = []
    current: dict[str, str] = {}

    for lineno, raw in enumerate(lines, start=1):
        line = raw.strip().lstrip("\ufeff")  # also drops a UTF-8 BOM
        if not line:
            if current:
                records.append(current)
                current = {}
            continue

        key, sep, value = line.partition(":")
        key = key.strip()
        if not sep or not key:
            log.warning("line %d skipped (expected 'key: value'): %r", lineno, line)
            continue
        if key in current:
            log.warning("line %d: duplicate key %r in one record, keeping the last value", lineno, key)
        current[key] = value.strip()

    if current:
        records.append(current)
    return records


def read_records(path: Path, encoding: str = "utf-8") -> list[dict[str, str]]:
    with path.open(encoding=encoding) as f:
        return parse_records(f)


def validate_sheet_name(name: str) -> str:
    if not name or len(name) > MAX_SHEET_NAME_LEN:
        raise ValueError(f"sheet name must be 1-{MAX_SHEET_NAME_LEN} characters")
    bad = INVALID_SHEET_CHARS.intersection(name)
    if bad:
        raise ValueError(f"sheet name contains invalid characters: {''.join(sorted(bad))}")
    return name


def write_excel(records: list[dict[str, str]], out: Path, sheet: str = DEFAULT_SHEET,
                rtl: bool = False) -> int:
    
    df = pd.DataFrame(records) 
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet)
        ws = writer.sheets[sheet]

        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for idx, column in enumerate(df.columns, start=1):
            longest = max([len(str(column))] + [len(str(v)) for v in df[column].dropna()])
            ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, 60)
        if rtl:
            ws.sheet_view.rightToLeft = True
    return len(df)


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Convert a text file of key: value records into an Excel file.")
    p.add_argument("input", nargs="?", default=DEFAULT_INPUT, type=Path,
                   help=f"input text file (default: {DEFAULT_INPUT})")
    p.add_argument("output", nargs="?", default=DEFAULT_OUTPUT, type=Path,
                   help=f"output .xlsx file (default: {DEFAULT_OUTPUT})")
    p.add_argument("--sheet", default=DEFAULT_SHEET, help=f"sheet name (default: {DEFAULT_SHEET})")
    p.add_argument("--encoding", default="utf-8", help="input file encoding (default: utf-8)")
    p.add_argument("--rtl", action="store_true", help="right-to-left sheet (for Persian/Arabic)")
    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    try:
        sheet = validate_sheet_name(args.sheet)
        records = read_records(args.input, args.encoding)
        if not records:
            log.error("no records found in %s", args.input)
            return 1
        count = write_excel(records, args.output, sheet, args.rtl)
    except FileNotFoundError:
        log.error("input file not found: %s", args.input)
        return 1
    except PermissionError:
        log.error("cannot write %s (is it open in Excel?)", args.output)
        return 1
    except UnicodeDecodeError:
        log.error("could not decode %s as %s; try --encoding", args.input, args.encoding)
        return 1
    except ValueError as e:
        log.error("%s", e)
        return 1

    log.info("saved %d record(s) to %s", count, args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
